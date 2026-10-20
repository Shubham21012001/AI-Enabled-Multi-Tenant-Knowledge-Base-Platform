"""Unit tests for extraction, cleaning, chunking, embeddings and prompt rules."""
import io

import pytest
from pypdf import PdfWriter

from app.core.config import settings
from app.repositories.chunk_repository import RetrievedChunk
from app.services import rag_service
from app.services.ai_provider import FakeProvider
from app.services.chunking import chunk_text
from app.services.extraction import ExtractionError, clean_text, extract_text
from tests.test_documents import make_docx


def _pdf_with_text(text: str) -> bytes:
    """Build a minimal one-page PDF containing `text` (no extra dependencies)."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(content) + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


# --- Extraction ---

def test_extract_txt():
    assert extract_text("Hello  world\r\n\r\n\r\n\r\nBye".encode(), "txt") == "Hello world\n\nBye"


def test_extract_txt_cp1252_fallback():
    assert "café" in extract_text("café menu".encode("cp1252"), "txt")


def test_extract_docx_including_tables():
    from docx import Document

    d = Document()
    d.add_paragraph("Holiday policy")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Annual leave"
    table.rows[0].cells[1].text = "24 days"
    buf = io.BytesIO()
    d.save(buf)
    text = extract_text(buf.getvalue(), "docx")
    assert "Holiday policy" in text and "Annual leave | 24 days" in text


def test_extract_pdf():
    text = extract_text(_pdf_with_text("Employees receive 24 days of leave"), "pdf")
    assert "24 days of leave" in text


def test_extract_blank_pdf_raises():
    w = PdfWriter()
    w.add_blank_page(612, 792)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(ExtractionError):
        extract_text(buf.getvalue(), "pdf")


def test_extract_docx_roundtrip():
    assert "paid sick leave" in extract_text(make_docx("Twelve days of paid sick leave."), "docx")


# --- Cleaning ---

def test_clean_text_rules():
    raw = "ﬁnal  report\x00\nemploy-\nees get\t\tleave\n\n\n\n\nEnd"
    assert clean_text(raw) == "final report\nemployees get leave\n\nEnd"


# --- Chunking ---

def test_short_text_single_chunk():
    assert chunk_text("short text") == ["short text"]
    assert chunk_text("   ") == []


def test_chunk_size_and_overlap():
    sentences = [f"Sentence number {i} talks about policy item {i}." for i in range(200)]
    text = " ".join(sentences)
    chunks = chunk_text(text, chunk_size=1000, overlap=200)
    assert len(chunks) > 1
    assert all(len(c) <= 1000 for c in chunks)
    # Consecutive chunks share text (the overlap).
    for a, b in zip(chunks, chunks[1:]):
        tail = a[-100:]
        assert any(tail[i : i + 30] in b for i in range(0, 70, 10))
    # Nothing is lost: every sentence appears in some chunk.
    joined = " ".join(chunks)
    assert all(s in joined for s in sentences)


def test_chunks_prefer_sentence_boundaries():
    text = ". ".join(["word " * 30] * 20)
    for c in chunk_text(text, chunk_size=1000, overlap=200)[:-1]:
        assert c.endswith(".")


def test_chunking_handles_text_without_spaces():
    chunks = chunk_text("x" * 2500, chunk_size=1000, overlap=200)
    assert all(len(c) <= 1000 for c in chunks) and len(chunks) == 3


def test_invalid_overlap():
    with pytest.raises(ValueError):
        chunk_text("abc " * 500, chunk_size=100, overlap=100)


# --- Embeddings (offline provider) ---

def test_embedding_dimension_and_similarity():
    import numpy as np

    p = FakeProvider()
    a, b, c = p.embed(["annual leave days policy", "how many annual leave days", "rocket engine thrust"])
    assert len(a) == settings.EMBEDDING_DIM == 1536
    assert np.dot(a, b) > np.dot(a, c)


# --- Prompt / grounding ---

def _chunk(i: int, text: str) -> RetrievedChunk:
    import uuid

    return RetrievedChunk(uuid.uuid4(), uuid.uuid4(), f"doc{i}.pdf", i, text, 0.8)


def test_system_prompt_enforces_grounding():
    p = rag_service.SYSTEM_PROMPT
    assert "ONLY" in p and "outside" in p and "invent" in p
    assert settings.NOT_FOUND_MESSAGE in p


def test_user_prompt_contains_numbered_sources():
    prompt = rag_service.build_user_prompt("What is X?", [_chunk(0, "X is 1"), _chunk(1, "Y is 2")])
    assert "[Source 1]" in prompt and "[Source 2]" in prompt and prompt.endswith("QUESTION: What is X?")


def test_sources_follow_citations():
    chunks = [_chunk(0, "a"), _chunk(1, "b"), _chunk(2, "c")]
    sources = rag_service.build_sources("The answer is b [2].", chunks)
    assert [s["chunk_index"] for s in sources] == [1]
    assert {"document_id", "document_name", "chunk_index"} <= set(sources[0])
    # No citation -> all retrieved chunks
    assert len(rag_service.build_sources("Answer without refs", chunks)) == 3
    # Not found -> no sources
    assert rag_service.build_sources(settings.NOT_FOUND_MESSAGE, chunks) == []


def test_not_found_detection():
    assert rag_service.is_not_found(settings.NOT_FOUND_MESSAGE)
    assert rag_service.is_not_found(f'"{settings.NOT_FOUND_MESSAGE}"')
    assert not rag_service.is_not_found("Employees get 24 days of leave [1].")


def test_openai_error_classification():
    """Quota/auth errors are permanent; temporary rate limits and 5xx are retryable."""
    import httpx
    from openai import AuthenticationError, InternalServerError, RateLimitError

    from app.services.ai_provider import AIProviderUnavailable, _translate_openai_error

    def resp(status):
        return httpx.Response(status, request=httpx.Request("POST", "https://api.openai.com/v1/embeddings"))

    quota = RateLimitError("x", response=resp(429), body={"type": "insufficient_quota", "code": "credit_balance_exhausted"})
    slow = RateLimitError("x", response=resp(429), body={"type": "requests", "code": "rate_limit_exceeded"})
    auth = AuthenticationError("x", response=resp(401), body={"code": "invalid_api_key"})
    down = InternalServerError("x", response=resp(500), body=None)

    assert isinstance(_translate_openai_error(quota, "E"), AIProviderUnavailable)
    assert _translate_openai_error(quota, "E").reason == "insufficient_quota"
    assert isinstance(_translate_openai_error(auth, "E"), AIProviderUnavailable)
    assert not isinstance(_translate_openai_error(slow, "E"), AIProviderUnavailable)
    assert not isinstance(_translate_openai_error(down, "E"), AIProviderUnavailable)
