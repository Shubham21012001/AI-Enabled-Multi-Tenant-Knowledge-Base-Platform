"""Step 1 + 2 of the RAG pipeline: extract raw text, then clean it."""
import io
import re
import unicodedata

from docx import Document as DocxDocument
from pypdf import PdfReader


class ExtractionError(Exception):
    """The file is unreadable or contains no extractable text. Not retryable."""


def extract_text(data: bytes, file_type: str) -> str:
    try:
        if file_type == "pdf":
            text = _extract_pdf(data)
        elif file_type == "docx":
            text = _extract_docx(data)
        elif file_type == "txt":
            text = _extract_txt(data)
        else:
            raise ExtractionError(f"Unsupported file type: {file_type}")
    except ExtractionError:
        raise
    except Exception as exc:  # corrupted / encrypted files
        raise ExtractionError(f"Could not read {file_type.upper()} file") from exc
    text = clean_text(text)
    if not text:
        raise ExtractionError(
            "No text found in this document. Scanned/image-only PDFs are not supported yet."
        )
    return text


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ExtractionError("Password-protected PDFs are not supported") from exc
    return "\n\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    # Tables often hold the important facts in policies (leave days, prices...).
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def clean_text(text: str) -> str:
    """Normalise extracted text without destroying paragraph structure.

    - Unicode NFKC (turns ligatures like 'ﬁ' into 'fi')
    - remove control characters
    - re-join words hyphenated across PDF line breaks ("employ-\\nees")
    - collapse runs of spaces/tabs, and 3+ newlines into a paragraph break
    """
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
